#!/usr/bin/env python3
"""Baut die Mapping-Arbeitsmappe für eine Ablösestrecke.

Erzeugt aus einer Ist-Spaltenliste die Excel-Mappe für die Mapping-Sitzung: fünf
Entscheidungen vor der Feldliste, alle Spalten in Originalreihenfolge mit markierten
Fallstricken, Wertelisten und die Beschaffungsliste.

    # Leere Mappe für eine neue Strecke
    python3 werkzeuge/mappe_bauen.py spalten.txt -o Mapping_Strecke.xlsx --titel "§ 14a"

    # Mit Zuordnungsvorschlägen aus schema_attribute.py
    python3 werkzeuge/mappe_bauen.py spalten.txt -o Mapping.xlsx \
        --vorschlaege vorschlag.csv --titel "§ 14a"

Die fünf Entscheidungen sind bei jeder Ablösung dieselben — sie stehen deshalb fest in
der Vorlage. Streckenspezifisches kommt über --hinweis dazu.
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

sys.path.insert(0, str(Path(__file__).resolve().parent))
from spaltenanalyse import TEXT_MUSTER, befunde, gruppiere, lies

F = "Arial"
INK, HEADBG, PETROL, GREY = "1A1A1A", "1F3B42", "0E5A69", "5A6A73"
FUELL, REF, KRIT, OK, GRUPPE, WARN = "FFF2CC", "F2F2F2", "F8D7D2", "DEEBE4", "D9E2E5", "F6EBD9"
thin = Side(style="thin", color="BFBFBF")
box = Border(left=thin, right=thin, top=thin, bottom=thin)

# Diese Entscheidungen fallen bei jeder Ablösung an - deshalb fest in der Vorlage.
ENTSCHEIDUNGEN = [
    ("1", "Welche Spalten braucht das Zielsystem inhaltlich?",
     "Anzeigetexte und Bestätigungsfelder müssen oft nur formal in der Datei stehen. Werden "
     "sie beim Import verworfen, entfällt die Zuordnung für sie komplett — das kann einen "
     "erheblichen Teil der Feldarbeit sparen.",
     "Verarbeitet das Zielsystem alle Spalten oder nur einen Teil? Bei welchen zählt der Inhalt?"),
    ("2", "Erkennt das Zielsystem die Spalten an Position oder Name?",
     "Bei doppelten Spaltennamen ist ein Mapping über Namen unmöglich. Die neue Erzeugung "
     "muss dann dieselbe Spalte an derselben Stelle ausgeben.",
     "Position bestätigen lassen. Falls der Name gelesen wird: Wie wird heute mit Duplikaten "
     "umgegangen?"),
    ("3", "Was tritt an die Stelle der Altsystem-Kennung?",
     "Wird das Altsystem abgelöst, verliert sein Kennungsfeld die Quelle. Drei Wege: "
     "Nummernkreis fortführen, neue Kennung im Zielsystem akzeptieren, Umsetzungstabelle "
     "führen. Nur der erste lässt die Zielseite unberührt.",
     "Einheitlich über alle Strecken entscheiden. Und: Was wird aus Bestandsdaten mit alter "
     "Kennung?"),
    ("4", "Bedingte Pflichtfelder festlegen",
     "Wo ein Feld die Sichtbarkeit anderer Blöcke steuert, ist die Datei dünn besetzt. Eine "
     "Prüfung, die alle Felder immer verlangt, hält jeden Vorgang zurück.",
     "Welches Feld ist der Diskriminator? Welche Blöcke sind je Ausprägung Pflicht?"),
    ("5", "Kopplung von Formular und Schnittstelle",
     "Bildet der Export die Eingabemaske ab, ist die Formularstruktur Teil der "
     "Schnittstellenspezifikation — wer Felder umbenennt oder umsortiert, verändert die Datei.",
     "Wird das akzeptiert, oder soll die Kopplung bei der Ablösung aufgelöst werden? "
     "Letzteres bedeutet eine Anpassung im Zielsystem."),
]

BESCHAFFEN = [
    ("1", "Produktive Originaldatei mit Inhalten",
     "Klärt Zeichenkodierung, Trennzeichen, Zeilenende und Datumsformate — und ist zugleich "
     "der Referenzsatz für den Abnahmevergleich."),
    ("2", "Zugang zum bestehenden Erzeugungscode",
     "Enthält Mapping, Transformationen und Prüfregeln vollständig, inklusive der Sonderfälle."),
    ("3", "Zielpfad und Schreibrechte",
     "Dorthin schreibt der neue Lauf. Braucht ein technisches Konto."),
    ("4", "Auskunft: Wie wird ein fehlgeschlagener Import heute bemerkt?",
     "Eine Dateischnittstelle hat keine Quittung. Ohne Abgleich bleibt unentdeckt, dass die "
     "Datei gar nicht abgeholt wurde."),
    ("5", "Mengengerüst aus dem Protokoll der heutigen Erzeugung",
     "Entscheidet über Takt und Verfahren. Muss nicht geschätzt werden."),
    ("6", "Abschalttermin des Altsystems",
     "Setzt den Rahmen für die Planung — der Endtermin kommt von außen."),
]


def head(
    ws: Worksheet,
    row: int,
    labels: list[str],
    widths: list[float],
    fills: list[str] | None = None,
    h: float = 38,
) -> None:
    """Kopfzeile eines Blattes setzen und darunter einfrieren.

    Args:
        ws: Das Arbeitsblatt.
        row: Zeilennummer der Kopfzeile.
        labels: Spaltenüberschriften.
        widths: Spaltenbreiten in Zeichen, in derselben Reihenfolge.
        fills: Hintergrundfarben je Spalte; `None` färbt alle einheitlich.
        h: Höhe der Kopfzeile in Punkt.
    """
    for i, t in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=t)
        c.font = Font(name=F, size=9, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=(fills[i - 1] if fills else HEADBG))
        c.alignment = Alignment(vertical="center", wrap_text=True)
        c.border = box
    ws.row_dimensions[row].height = h
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def cell(
    ws: Worksheet,
    r: int,
    c: int,
    v: str | int | None = None,
    fill: str | None = None,
    bold: bool = False,
    size: float = 10,
    italic: bool = False,
    color: str = INK,
) -> Cell:
    """Eine Zelle setzen und einheitlich formatieren.

    Args:
        ws: Das Arbeitsblatt.
        r: Zeile, 1-basiert.
        c: Spalte, 1-basiert.
        v: Inhalt; `None` lässt die Zelle leer, aber formatiert sie.
        fill: Hintergrundfarbe als Hexwert ohne `#`.
        bold: Fett setzen.
        size: Schriftgröße in Punkt.
        italic: Kursiv setzen.
        color: Schriftfarbe als Hexwert ohne `#`.

    Returns:
        Die gesetzte Zelle.
    """
    x = ws.cell(row=r, column=c, value=v)
    x.font = Font(name=F, size=size, bold=bold, italic=italic, color=color)
    x.alignment = Alignment(vertical="top", wrap_text=True)
    x.border = box
    if fill:
        x.fill = PatternFill("solid", fgColor=fill)
    return x


def titel(
    ws: Worksheet,
    t: str,
    sub: str,
    hin: str | None = None,
    hin2: str | None = None,
) -> None:
    """Titelblock über die Kopfzeile setzen.

    Args:
        ws: Das Arbeitsblatt.
        t: Überschrift.
        sub: Erläuterung darunter.
        hin: Hervorgehobener Hinweis, meist der Fallstrick dieses Blattes.
        hin2: Nachgestellter Hinweis in kleiner Schrift, etwa ein Werkzeugaufruf.
    """
    ws.sheet_view.showGridLines = False
    ws.cell(row=1, column=1, value=t).font = Font(name=F, size=14, bold=True, color=INK)
    ws.cell(row=2, column=1, value=sub).font = Font(name=F, size=9.5, color=GREY)
    if hin:
        ws.cell(row=3, column=1, value=hin).font = Font(name=F, size=9.5, bold=True, color="8A5510")
    if hin2:
        ws.cell(row=4, column=1, value=hin2).font = Font(name=F, size=9, italic=True, color=GREY)


def lies_vorschlaege(pfad: str) -> dict[int, list[str]]:
    """Zuordnungsvorschläge aus `schema_attribute.py` einlesen.

    Args:
        pfad: CSV-Datei, wie `schema_attribute.py --spalten` sie schreibt.

    Returns:
        Je Spaltenposition die Vorschläge, jeder mit seiner Güte in Klammern.
        Positionen ohne Vorschlag fehlen im Ergebnis.
    """
    ergebnis: dict[int, list[str]] = {}
    with open(pfad, encoding="utf-8-sig", newline="") as f:
        for zeile in csv.DictReader(f, delimiter=";"):
            try:
                pos = int(zeile.get("position", ""))
            except ValueError:
                continue
            treffer = []
            for n in ("1", "2", "3"):
                wert = (zeile.get(f"vorschlag_{n}") or "").strip()
                if wert:
                    guete = (zeile.get(f"guete_{n}") or "").strip()
                    treffer.append(f"{wert}" + (f" ({guete})" if guete else ""))
            if treffer:
                ergebnis[pos] = treffer
    return ergebnis


def baue(
    namen: list[str],
    ziel: str,
    titel_text: str,
    vorschlaege: dict[int, list[str]],
    hinweis: str | None,
) -> None:
    """Die vollständige Mapping-Arbeitsmappe schreiben.

    Vier Blätter: Entscheidungen, Feldmapping, Wertelisten, Beschaffungsliste. Die
    Spaltenreihenfolge des Ist-Exports bleibt unverändert — sie ist bei doppelten
    Spaltennamen die einzige verlässliche Zuordnung.

    Args:
        namen: Spaltennamen des Ist-Exports in Originalreihenfolge.
        ziel: Pfad der zu schreibenden .xlsx-Datei.
        titel_text: Name der Strecke, erscheint in den Überschriften.
        vorschlaege: Zuordnungsvorschläge je Position; leer lässt die Spalte weg.
        hinweis: Streckenspezifischer Hinweis auf Blatt 1.
    """
    b = befunde(namen)
    gruppen_von = {n: g for g, ms in gruppiere(namen).items() for n in ms}
    wb = Workbook()

    # ---------------------------------------------------- Entscheidungen
    ws = wb.active
    ws.title = "1 Entscheidungen"
    titel(ws, f"Mapping-Sitzung {titel_text}",
          "Diese fünf Punkte vor der Feldliste — sie bestimmen, wie viel Feldarbeit anfällt.",
          hinweis
          or "Punkt 1 kann einen erheblichen Teil der Arbeit sparen. Deshalb steht er vorn.")
    head(ws, 6, ["Nr", "Punkt", "Warum zuerst", "Zu entscheiden", "Entscheidung", "Wer", "Bis"],
         [5, 30, 44, 34, 30, 14, 10], [HEADBG] * 4 + [FUELL] * 3, h=40)
    r = 7
    for nr, punkt, warum, was in ENTSCHEIDUNGEN:
        cell(ws, r, 1, nr, REF, bold=True, size=9)
        cell(ws, r, 2, punkt, KRIT if nr in ("1", "3") else REF, bold=True)
        cell(ws, r, 3, warum, REF, size=9.5, italic=True, color="4A5A63")
        cell(ws, r, 4, was, REF, size=9.5)
        for c_ in (5, 6, 7):
            cell(ws, r, c_, None, FUELL)
        ws.row_dimensions[r].height = 76
        r += 1

    # ---------------------------------------------------- Feldmapping
    ws = wb.create_sheet("2 Feldmapping")
    hat_vorschlaege = bool(vorschlaege)
    titel(ws, f"Feldmapping {titel_text} — {len(namen)} Spalten",
          "Reihenfolge und Schreibweise exakt wie im heutigen Export. Position ist verbindlich.",
          ("Spalte „Vorschlag“ stammt aus dem Namensabgleich und ist ungeprüft — "
           "Güte unter 0,75 anschauen."
           if hat_vorschlaege else
           "Die Spalte „Quelle im Zielsystem“ ist leer: Es liegt noch kein Mapping-Vorschlag vor."),
          "Vorschläge erzeugen: werkzeuge/schema_attribute.py --manifest … --spalten …")
    spalten = ["Pos", "Spaltenname", "Block", "Art", "Fallstrick / Hinweis"]
    breiten = [5, 32, 16, 10, 38]
    fuellungen = [HEADBG] * 5
    if hat_vorschlaege:
        spalten.append("Vorschlag (ungeprüft)")
        breiten.append(30)
        fuellungen.append(WARN)
    spalten += ["Wird inhaltlich gebraucht?", "Quelle im Zielsystem", "Transformation", "Wer"]
    breiten += [20, 28, 24, 10]
    fuellungen += [FUELL] * 4
    head(ws, 6, spalten, breiten, fuellungen, h=40)

    r = 7
    letzter = None
    for i, name in enumerate(namen, 1):
        t = name.strip()
        block = gruppen_von.get(name, "")
        if block != letzter and block and block != "(einzeln)":
            cell(ws, r, 2, block, GRUPPE, bold=True, size=9.5, color=PETROL)
            for c_ in [1, *range(3, len(spalten) + 1)]:
                cell(ws, r, c_, None, GRUPPE)
            ws.row_dimensions[r].height = 16
            r += 1
            letzter = block
        art = "Anzeige?" if TEXT_MUSTER.match(t) else "Daten"
        hinweise = []
        if name in b["mehrfach"]:
            hinweise.append("MEHRFACH — Position " + ", ".join(map(str, b["mehrfach"][name])))
        if name != t:
            hinweise.append("LEERZEICHEN AM RAND im Spaltennamen")
        if re.search(r"[./§%&]", t):
            hinweise.append("Sonderzeichen im Namen")
        if " " in t:
            hinweise.append("Leerzeichen im Namen")
        auffaellig = name in b["mehrfach"] or name != t
        f = KRIT if auffaellig else (WARN if art == "Anzeige?" else REF)
        cell(ws, r, 1, i, REF, size=9)
        cell(ws, r, 2, name, f, bold=True, size=9.5)
        cell(ws, r, 3, block, REF, size=9)
        cell(ws, r, 4, art, WARN if art == "Anzeige?" else REF, size=9)
        cell(ws, r, 5, " · ".join(hinweise) or None, REF, size=8.5, italic=True,
             color="8E3324" if auffaellig else "4A5A63")
        c_ = 6
        if hat_vorschlaege:
            cell(ws, r, c_, "\n".join(vorschlaege.get(i, [])) or None, WARN, size=9)
            c_ += 1
        for _ in range(4):
            cell(ws, r, c_, None, FUELL, size=9.5)
            c_ += 1
        ws.row_dimensions[r].height = 26
        r += 1

    sp_gebraucht = get_column_letter(6 + (1 if hat_vorschlaege else 0))
    dv = DataValidation(type="list", formula1='"ja,nein,leer mitliefern,klären"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{sp_gebraucht}7:{sp_gebraucht}{r}")

    # ---------------------------------------------------- Wertelisten
    ws = wb.create_sheet("3 Wertelisten")
    titel(ws, f"Wertelisten {titel_text}",
          "Für jedes Auswahlfeld: welchen Schlüssel erwartet das Zielsystem?",
          "Ausprägungen aus einer produktiven Originaldatei ablesen, nicht raten.")
    head(ws, 5, ["Feld", "Klartext / Ausprägung", "Wert im Quellsystem", "Schlüssel im Zielsystem",
                 "Bemerkung"], [26, 34, 26, 24, 32],
         [HEADBG, HEADBG, FUELL, FUELL, HEADBG])
    r = 6
    for _ in range(30):
        for c_ in range(1, 6):
            cell(ws, r, c_, None, FUELL if c_ in (3, 4) else REF, size=9.5)
        r += 1

    # ---------------------------------------------------- Beschaffen
    ws = wb.create_sheet("4 Noch zu beschaffen")
    titel(ws, "Was für die Umsetzung fehlt", "Unabhängig vom Mapping.")
    head(ws, 5, ["Nr", "Was", "Warum", "Wer", "Bis"], [5, 38, 46, 20, 12],
         [HEADBG, HEADBG, HEADBG, FUELL, FUELL])
    r = 6
    for nr, was, warum in BESCHAFFEN:
        cell(ws, r, 1, nr, REF, bold=True, size=9)
        cell(ws, r, 2, was, REF, bold=True, size=9.5)
        cell(ws, r, 3, warum, REF, size=9, italic=True, color="4A5A63")
        cell(ws, r, 4, None, FUELL)
        cell(ws, r, 5, None, FUELL)
        ws.row_dimensions[r].height = 44
        r += 1

    wb.save(ziel)


def main(argv: list[str] | None = None) -> int:
    """Einstiegspunkt für den Aufruf über die Kommandozeile.

    Args:
        argv: Argumente; `None` nimmt die der Kommandozeile.

    Returns:
        0 — das Werkzeug erzeugt eine Arbeitsvorlage und bewertet nichts.
    """
    p = argparse.ArgumentParser(description=__doc__,
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("spalten", help="Datei mit den Spaltennamen des Ist-Exports")
    p.add_argument("-o", "--ausgabe", required=True, help="Zieldatei (.xlsx)")
    p.add_argument("--titel", default="Strecke")
    p.add_argument("--vorschlaege", help="CSV aus schema_attribute.py")
    p.add_argument("--hinweis", help="streckenspezifischer Hinweis auf Blatt 1")
    a = p.parse_args(argv)

    namen = lies(a.spalten)
    vorschlaege = lies_vorschlaege(a.vorschlaege) if a.vorschlaege else {}
    baue(namen, a.ausgabe, a.titel, vorschlaege, a.hinweis)

    b = befunde(namen)
    print(f"geschrieben: {a.ausgabe}", file=sys.stderr)
    print(f"  {len(namen)} Spalten · {len(b['anzeigetexte'])} mutmaßliche Anzeigetexte "
          f"· {len(b['mehrfach'])} doppelte Namen "
          f"· {len(b['randleerzeichen'])} mit Randleerzeichen",
          file=sys.stderr)
    if vorschlaege:
        print(f"  {len(vorschlaege)} Spalten mit Zuordnungsvorschlag", file=sys.stderr)
    return 0


if __name__ == "__main__":
    sys.exit(main())
